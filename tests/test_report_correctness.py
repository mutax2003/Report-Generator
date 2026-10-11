"""Regression tests for report-output correctness (escaping, sandbox limits, formatting)."""

from __future__ import annotations

import io
import sys
import unittest
import zipfile
from pathlib import Path

import openpyxl
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from engine import ReportEngine  # noqa: E402

META = {"prepared_by": "QA", "date_of_issue": "2026-10-07", "report_phase": "Phase 2"}
LAB_HDR = ["Analyte", "Result", "Unit", "Criteria", "Exceedance"]


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _template(lines: list[str], *, lab_table: bool = False) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if lab_table:
        table = doc.add_table(rows=3, cols=3)
        table.rows[0].cells[0].merge(table.rows[0].cells[2])
        table.rows[0].cells[0].text = "{%tr for item in lab_results %}"
        table.rows[1].cells[0].text = "{{ item.analyte }}"
        table.rows[1].cells[1].text = "{{ item.result_display }}"
        table.rows[1].cells[2].text = "{{ item.exceedance_flag }}"
        table.rows[2].cells[0].merge(table.rows[2].cells[2])
        table.rows[2].cells[0].text = "{%tr endfor %}"
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _document_xml(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


def _plain_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


class TestXmlEscaping(unittest.TestCase):
    """Defect 1: Excel text must be XML-escaped, never parsed as WordprocessingML."""

    def _render(self, client: str, notes: str = "plain") -> bytes:
        xb = _xlsx(
            {
                "ProjectData": [
                    ["site_name", "client_name", "project_number", "notes"],
                    ["Smith & Sons Site", client, "P-1", notes],
                ],
                "LabResults": [
                    LAB_HDR,
                    ["Benzene", "<0.005", "mg/kg", 0.01, "N"],
                    ["Toluene & Xylene", ">10", "mg/kg", 0.5, "Y"],
                ],
            }
        )
        tb = _template(
            ["Site: {{ site_name }}", "Client: {{ client_name }}", "Notes: {{ notes }}"],
            lab_table=True,
        )
        docx_bytes, _w, _ctx, _rec = ReportEngine(xb, tb).render(META)
        return docx_bytes

    def test_lt_gt_amp_quotes_survive(self) -> None:
        docx_bytes = self._render('O\'Brien "Holdings" <Ltd> & Co')
        text = _plain_text(docx_bytes)
        self.assertIn("Smith & Sons Site", text)
        self.assertIn('Client: O\'Brien "Holdings" <Ltd> & Co', text)
        self.assertIn("<0.005", text)
        self.assertIn("Toluene & Xylene", text)
        xml = _document_xml(docx_bytes)
        self.assertIn("&lt;0.005", xml)
        # Exceedance value goes through RichText (escapes its own text).
        self.assertIn("&gt;10</w:t>", xml)
        self.assertIn("Smith &amp; Sons", xml)

    def test_cell_cannot_inject_wordprocessingml(self) -> None:
        docx_bytes = self._render("</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>INJECTED")
        xml = _document_xml(docx_bytes)
        self.assertNotIn("<w:t>INJECTED", xml)
        self.assertIn("&lt;/w:t&gt;", xml)
        self.assertIn("</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>INJECTED", _plain_text(docx_bytes))

    def test_multiline_text_becomes_line_breaks(self) -> None:
        docx_bytes = self._render("C", notes="Line one & more\nLine <two>")
        xml = _document_xml(docx_bytes)
        self.assertIn("<w:br/>", xml)
        text = _plain_text(docx_bytes)
        self.assertIn("Line one & more", text)
        self.assertIn("Line <two>", text)

    def test_richtext_exceedance_still_renders_as_runs(self) -> None:
        docx_bytes = self._render("C")
        xml = _document_xml(docx_bytes)
        # >10 exceeds 0.5 (flag Y) → red bold RichText run, not escaped markup.
        self.assertIn('<w:color w:val="FF0000"/>', xml)
        self.assertNotIn("&lt;w:r&gt;", xml)

    def test_appendix_render_escapes(self) -> None:
        from appendix_generator import _render_appendix_docx

        tb = _template(["Client: {{ client_name }}", "Missing: {{ not_supplied }}"])
        out = _render_appendix_docx(tb, {"client_name": "A & B <Ltd>"})
        self.assertIn("A &amp; B &lt;Ltd&gt;", _document_xml(out))
        self.assertIn("Client: A & B <Ltd>", _plain_text(out))


class TestSandboxLimits(unittest.TestCase):
    """Defect 2: template expressions cannot exhaust memory/CPU."""

    def _render_expr(self, expr: str) -> None:
        xb = _xlsx({"ProjectData": [["site_name"], ["S"]], "LabResults": [LAB_HDR]})
        ReportEngine(xb, _template([expr])).render(META)

    def test_huge_string_repetition_blocked(self) -> None:
        with self.assertRaises(ValueError) as cm:
            self._render_expr("{{ ('a' * 10**7) | length }}")
        self.assertIn("repetition", str(cm.exception))

    def test_nested_range_budget_blocked(self) -> None:
        with self.assertRaises(ValueError) as cm:
            self._render_expr(
                "{% for i in range(10000) %}{% for j in range(1000) %}{% endfor %}{% endfor %}"
            )
        self.assertIn("range", str(cm.exception))

    def test_huge_exponent_blocked(self) -> None:
        with self.assertRaises(ValueError):
            self._render_expr("{{ 10 ** 100000 }}")

    def test_normal_expressions_still_work(self) -> None:
        from report_jinja import make_report_jinja_env

        env = make_report_jinja_env()
        tpl = env.from_string("{{ 'ab' * 3 }}|{{ [1] + [2] }}|{{ 2 ** 10 }}|{{ range(3) | list }}")
        self.assertEqual(tpl.render(), "ababab|[1, 2]|1024|[0, 1, 2]")

    def test_range_budget_is_per_environment(self) -> None:
        from report_jinja import MAX_RANGE_BUDGET, make_report_jinja_env

        src = "{% for i in range(" + str(MAX_RANGE_BUDGET) + ") %}{% endfor %}ok"
        self.assertEqual(make_report_jinja_env().from_string(src).render(), "ok")
        self.assertEqual(make_report_jinja_env().from_string(src).render(), "ok")


class TestNegativeNumbers(unittest.TestCase):
    """Defect 3: no formula-injection apostrophe in Word output."""

    def test_negative_values_render_without_apostrophe(self) -> None:
        xb = _xlsx(
            {
                "ProjectData": [
                    ["site_name", "longitude", "elev_text"],
                    ["Neg Site", -114.5, "-3.2"],
                ],
                "LabResults": [LAB_HDR, ["Redox", -0.3, "mV", 5, "N"]],
            }
        )
        tb = _template(["LON: {{ longitude }}", "ELEV: {{ elev_text }}"], lab_table=True)
        docx_bytes, _w, ctx, _rec = ReportEngine(xb, tb).render(META)
        text = _plain_text(docx_bytes)
        self.assertIn("LON: -114.5", text)
        self.assertIn("ELEV: -3.2", text)
        self.assertIn("Redox | -0.3", text)
        self.assertNotIn("'-", text)
        self.assertEqual(ctx["lab_results"][0]["result"], "-0.3")


class TestBatchIsolation(unittest.TestCase):
    """Defect 4: a site with no linked rows must not inherit another site's rows."""

    def _engine(self, lab_rows: list[list]) -> ReportEngine:
        xb = _xlsx(
            {
                "ProjectData": [
                    ["site_name", "client_name", "project_number"],
                    ["Site A", "C", "P-A"],
                    ["Site B", "C", "P-B"],
                ],
                "LabResults": lab_rows,
            }
        )
        return ReportEngine(xb, _template(["Site: {{ site_name }}"], lab_table=True))

    def test_batch_site_without_rows_gets_empty_table(self) -> None:
        engine = self._engine(
            [
                ["site_name"] + LAB_HDR,
                ["Site A", "ANALYTE_FROM_SITE_A", 1, "mg/kg", 5, "N"],
                ["Site A", "ANOTHER_A_ROW", 2, "mg/kg", 5, "N"],
            ]
        )
        batch = engine.render_batch(META)
        self.assertEqual(len(batch), 2)
        self.assertEqual(len(batch[0].context["lab_results"]), 2)
        self.assertEqual(batch[1].context["lab_results"], [])
        self.assertNotIn("ANALYTE_FROM_SITE_A", _plain_text(batch[1].docx_bytes))
        self.assertTrue(any("lab_results" in w for w in batch[1].warnings), batch[1].warnings)
        self.assertFalse(any("links rows to sites" in w for w in batch[0].warnings))

    def test_single_render_of_unmatched_row_is_empty(self) -> None:
        engine = self._engine(
            [["site_name"] + LAB_HDR, ["Site A", "ANALYTE_FROM_SITE_A", 1, "mg/kg", 5, "N"]]
        )
        _docx, warnings, ctx, _rec = engine.render(META, project_row_index=1)
        self.assertEqual(ctx["lab_results"], [])
        self.assertTrue(any("links rows to sites" in w for w in warnings), warnings)

    def test_sheet_without_link_column_keeps_all_rows(self) -> None:
        engine = self._engine([LAB_HDR, ["Benzene", 1, "mg/kg", 5, "N"]])
        batch = engine.render_batch(META)
        self.assertEqual([len(b.context["lab_results"]) for b in batch], [1, 1])

    def test_blank_link_column_keeps_all_rows(self) -> None:
        engine = self._engine([["site_name"] + LAB_HDR, [None, "Benzene", 1, "mg/kg", 5, "N"]])
        batch = engine.render_batch(META)
        self.assertEqual([len(b.context["lab_results"]) for b in batch], [1, 1])

    def test_indexed_and_linear_filters_agree_on_no_match(self) -> None:
        from engine import (
            _filter_records_for_project,
            _filter_records_for_project_indexed,
            _index_records_by_link_columns,
        )

        records = [{"site_name": "Site A", "analyte": "Benzene"}]
        project = {"site_name": "Site B"}
        indexes = _index_records_by_link_columns(records)
        self.assertEqual(_filter_records_for_project(records, project), [])
        self.assertEqual(_filter_records_for_project_indexed(records, project, indexes), [])


class TestValueFormatting(unittest.TestCase):
    """Defect 5: dates, integral floats and blank phrase cells render cleanly."""

    def test_cell_str_formats(self) -> None:
        import datetime as dt

        import pandas as pd

        from engine import _cell_str

        self.assertEqual(_cell_str(dt.datetime(2026, 5, 20)), "2026-05-20")
        self.assertEqual(_cell_str(pd.Timestamp("2026-05-20")), "2026-05-20")
        self.assertEqual(_cell_str(dt.datetime(2026, 5, 21, 13, 5)), "2026-05-21 13:05:00")
        self.assertEqual(_cell_str(dt.date(2026, 5, 22)), "2026-05-22")
        self.assertEqual(_cell_str(pd.NaT), "")
        self.assertEqual(_cell_str(float("nan")), "")
        self.assertEqual(_cell_str(710.0), "710")
        self.assertEqual(_cell_str(-3.0), "-3")
        self.assertEqual(_cell_str(0.5), "0.5")
        self.assertEqual(_cell_str(710), "710")

    def test_dates_integers_and_phrase_blanks_in_docx(self) -> None:
        import datetime as dt

        xb = _xlsx(
            {
                "ProjectData": [
                    ["site_name", "sample_date", "well_depth_m", "site_recon_intro_selected"],
                    ["Row2 Site", dt.datetime(2026, 5, 20), 710, "custom_blank"],
                    ["Row3 Site", dt.datetime(2026, 5, 21), None, "custom_blank"],
                ],
                "LabResults": [LAB_HDR, ["Benzene", 0.1, "mg/kg", 1.0, "N"]],
                "PhraseCatalog": [
                    ["phrase_key", "option_id", "text"],
                    ["site_recon_intro", "custom_blank", None],
                    [None, None, None],
                ],
            }
        )
        tb = _template(
            ["DATE: {{ sample_date }}", "DEPTH: {{ well_depth_m }}", "PHRASE: {{ site_recon_intro }}"],
            lab_table=True,
        )
        docx_bytes, _w, ctx, _rec = ReportEngine(xb, tb).render(META)
        text = _plain_text(docx_bytes)
        self.assertIn("DATE: 2026-05-20\n", text)
        self.assertIn("DEPTH: 710\n", text)
        self.assertNotIn("00:00:00", text)
        self.assertNotIn("710.0", text)
        self.assertNotIn("nan", text.lower())
        self.assertEqual(ctx["lab_results"][0]["criteria"], "1")

    def test_phrase_rows_skip_nan(self) -> None:
        import pandas as pd

        from phrase_resolver import phrase_rows_from_dataframe

        df = pd.DataFrame(
            [
                {"phrase_key": "k", "option_id": "a", "text": float("nan")},
                {"phrase_key": float("nan"), "option_id": "b", "text": "x"},
                {"phrase_key": "k", "option_id": "c", "text": "Real text"},
            ]
        )
        self.assertEqual(phrase_rows_from_dataframe(df), (("k", "c", "Real text"),))


class TestNumericTextPreserved(unittest.TestCase):
    """Defect 6: text cells keep their digits; numeric comparisons still work."""

    def test_text_cells_not_coerced(self) -> None:
        xb = _xlsx(
            {
                "ProjectData": [
                    ["site_name", "well_code", "conc_text"],
                    ["S", "007", "0.50"],
                ],
                "LabResults": [LAB_HDR, ["Chloride", "120", "mg/L", 250, "N"]],
            }
        )
        tb = _template(["WELLID: {{ well_code }}", "CONC: {{ conc_text }}"], lab_table=True)
        docx_bytes, _w, ctx, _rec = ReportEngine(xb, tb).render(META)
        text = _plain_text(docx_bytes)
        self.assertIn("WELLID: 007", text)
        self.assertIn("CONC: 0.50", text)
        self.assertEqual(ctx["lab_results"][0]["result"], "120")

    def test_lab_exceedance_comparison_on_text_and_numeric_cells(self) -> None:
        xb = _xlsx(
            {
                "ProjectData": [["site_name"], ["S"]],
                "LabResults": [
                    LAB_HDR,
                    ["TextOver", "120", "mg/L", 100, None],
                    ["TextUnder", "0.50", "mg/L", "1.0", None],
                    ["NumOver", 85, "mg/L", 50, None],
                    ["NonDetect", "<0.005", "mg/L", 0.01, None],
                ],
            }
        )
        ctx = ReportEngine(xb, _template(["x"], lab_table=True)).build_context(META)
        flags = {r["analyte"]: r["exceedance_flag"] for r in ctx["lab_results"]}
        self.assertEqual(
            flags, {"TextOver": "Yes", "TextUnder": "No", "NumOver": "Yes", "NonDetect": "No"}
        )
        self.assertEqual(ctx["lab_results"][1]["result"], "0.50")

    def test_groundwater_exceedance_summary_with_text_numbers(self) -> None:
        xb = _xlsx(
            {
                "ProjectData": [["site_name", "client_name"], ["GW Site", "C"]],
                "MonitoringWells": [["well_id"], ["MW-01"]],
                "WaterLevels": [["well_id", "depth_to_water_m"], ["MW-01", 3.0]],
                "GroundwaterLab": [
                    ["well_id"] + LAB_HDR,
                    ["MW-01", "Chloride", "120", "mg/L", 100, None],
                    ["MW-01", "Sodium", 85.0, "mg/L", 200, None],
                ],
            }
        )
        meta = dict(META, report_type="groundwater_monitoring", report_phase="Phase 1")
        ctx = ReportEngine(xb, _template(["{{ site_name }}"])).build_context(meta)
        rows = {r["analyte"]: r for r in ctx["groundwater_results"]}
        self.assertEqual(rows["Chloride"]["result"], "120")
        self.assertEqual(rows["Sodium"]["result"], "85")
        self.assertIn("Chloride", ctx["exceedance_summary"])
        self.assertNotIn("Sodium", ctx["exceedance_summary"])
        self.assertEqual(ctx["water_levels"][0]["depth_to_water_m"], "3")

    def test_dwda_parse_float_on_context_strings(self) -> None:
        import pandas as pd

        from compliance_helpers import parse_float
        from engine import _dataframe_to_records

        recs = _dataframe_to_records(
            pd.DataFrame([{"Volume": 1234.0, "Text Vol": "1,234", "Blank": None}], dtype=object)
        )
        self.assertEqual(recs[0]["volume"], "1234")
        self.assertEqual(parse_float(recs[0]["volume"]), 1234.0)
        self.assertEqual(parse_float(recs[0]["text_vol"]), 1234.0)
        self.assertIsNone(parse_float(recs[0]["blank"]))


class TestBlankLabRows(unittest.TestCase):
    """Defect 7: fully blank spreadsheet rows are not results."""

    def test_blank_lab_row_skipped(self) -> None:
        xb = _xlsx(
            {
                "ProjectData": [["site_name"], ["S8"]],
                "LabResults": [
                    LAB_HDR,
                    ["Benzene", 0.1, "mg/kg", 1, "N"],
                    [None] * 5,
                    ["Toluene", 0.2, "mg/kg", 1, "N"],
                ],
            }
        )
        ctx = ReportEngine(xb, _template(["x"], lab_table=True)).build_context(META)
        self.assertEqual([r["analyte"] for r in ctx["lab_results"]], ["Benzene", "Toluene"])
        self.assertEqual([r["exceedance_flag"] for r in ctx["lab_results"]], ["No", "No"])

    def test_blank_generic_table_row_skipped(self) -> None:
        import pandas as pd

        from engine import _dataframe_to_records

        df = pd.DataFrame(
            [{"apec_id": "APEC-1", "description": "Tank"}, {"apec_id": None, "description": None}],
            dtype=object,
        )
        self.assertEqual(_dataframe_to_records(df), [{"apec_id": "APEC-1", "description": "Tank"}])


if __name__ == "__main__":
    unittest.main()
